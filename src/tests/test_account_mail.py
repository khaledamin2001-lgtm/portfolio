#!/usr/bin/env python3
"""The account email job (src/jobs/run_account_mail.py) end to end with a fake Firebase and a fake mailer, on synthetic
data and throwaway keys: an opted-in account's package opens with the mail key, its documents are read with its own
token and opened with its own key, new heads-up items and the Thursday summary are emailed once each, the record of what
was sent is saved back encrypted, a second run sends nothing, and an account that did not opt in is never touched.
A second account connected its Gmail: a synthetic Thndr monthly statement (src/tests/fixtures/make_statement_pdf.py) is
"fetched" by a fake imap_fetch, sync.js posts it from the day after the account's tracking start (the opening rows typed
at sign-up are kept, the earlier statement rows are not added twice), everything is saved in one pinned Firestore commit
encrypted to that account's key, the summary email goes to its own address, the month-end report (Excel workbook, and
the PDF factsheet when Playwright is there; REQUIRE_PDF=1 insists on it) is emailed to it as attachments and stamped, a
second run does nothing new, and a wrong app password is emailed once and recorded without touching the portfolio.
Friends: a friend request is emailed once; each account's copy for its friends (shares/{owner}/to/{friend}) opens only
with the friend's key and carries the portfolio documents without the Thndr account or the Gmail login; the site owner's
verified account shares the MAIN portfolio (the engine's documents) while another account asking for that gets its own;
status/{uid}.job is written for the admin screen.
History import: an account created with "Build it from my Thndr emails" starts from its earliest monthly statement in
Gmail (Aug-26: holdings and cash as opening rows) and the later one (Sep-26: its rows, then its snapshot's one extra share
and the cash matched with labelled adjustments) is built on top; the Gmail import starts after the latest; only the
latest month's report is emailed, with one summary email; an account with no monthly statement yet is told once and
nothing in its portfolio changes.
    python3 src/tests/test_account_mail.py <synthetic export dir> <tools dir with node_modules/pdfjs-dist>   (exit 0 = all pass)"""
import os, sys, json, gzip, base64, hashlib, shutil, tempfile, subprocess, urllib.parse
from email.message import EmailMessage
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "jobs"))
import store, imap_fetch, run_account_mail as ram   # noqa: E402
from cryptography.hazmat.primitives.asymmetric import ec
from cryptography.hazmat.primitives import serialization, hashes
from cryptography.hazmat.primitives.kdf.hkdf import HKDF
from cryptography.hazmat.primitives.ciphers.aead import AESGCM

fails = 0
def check(name, ok, detail=""):
    global fails
    print(("PASS " if ok else "FAIL ") + name + (f" · {detail}" if detail and not ok else ""))
    fails += 0 if ok else 1

SYN, PDFJS_TOOLS = sys.argv[1], sys.argv[2]
tmp = tempfile.mkdtemp()
try:
    # the mail key = a throwaway stand-in for p/khaled/keys.json, opened with its setup key like the job does
    code = os.path.join(tmp, "code"); os.makedirs(os.path.join(code, "p", "khaled"))
    shutil.copytree(os.path.join(ROOT), os.path.join(code, "src"), ignore=shutil.ignore_patterns("__pycache__"))
    shutil.copytree(os.path.join(PDFJS_TOOLS, "node_modules"), os.path.join(code, "src", "tools", "node_modules"), symlinks=True)
    sec = os.path.join(tmp, "sec")
    subprocess.run([sys.executable, os.path.join(ROOT, "site", "make_keys.py"), os.path.join(code, "p", "khaled", "keys.json"), sec], check=True, capture_output=True)
    os.environ["SETUP_KEY"] = open(os.path.join(sec, "setup_key.txt")).read().strip()
    mail_pub = store.load_keys(os.path.join(code, "p", "khaled", "keys.json"))["pub"]
    # an account: its own key pair, its documents from the synthetic export sealed to it (as the browser stores them)
    acct = ec.generate_private_key(ec.SECP256R1())
    pk8 = base64.b64encode(acct.private_bytes(serialization.Encoding.DER, serialization.PrivateFormat.PKCS8, serialization.NoEncryption())).decode()
    apub = base64.b64encode(acct.public_key().public_bytes(serialization.Encoding.X962, serialization.PublicFormat.UncompressedPoint)).decode()
    UID, OTHER = "Uacct1", "Uother"
    DB, t = {}, [0]
    def stamp():
        t[0] += 1; return f"2026-09-24T19:00:{t[0]:02d}.000000Z"
    for c in ("portfolio", "ledger"):
        for f in os.listdir(os.path.join(SYN, c)):
            x = json.load(open(os.path.join(SYN, c, f))); data = x.get("data", x) if isinstance(x, dict) and isinstance(x.get("data"), dict) else x
            DB[f"users/{UID}/docs/{c}__{f[:-5]}"] = {"blob": store.encode_doc({"pub": apub}, f[:-5], 1, data, "2026-09-20T00:00:00Z").decode().strip(), "updateTime": stamp()}
    DB[f"users/{OTHER}/docs/portfolio__settings"] = {"blob": "x", "updateTime": stamp()}
    # the package, sealed with the mail label to the mail key (what the site's Email updates switch writes)
    def seal_mail(obj):
        site = ec.EllipticCurvePublicKey.from_encoded_point(ec.SECP256R1(), base64.b64decode(mail_pub))
        eph = ec.generate_private_key(ec.SECP256R1())
        epk = eph.public_key().public_bytes(serialization.Encoding.X962, serialization.PublicFormat.UncompressedPoint)
        key = HKDF(hashes.SHA256(), 32, epk, b"portfolio-mail-v1").derive(eph.exchange(ec.ECDH(), site))
        iv = os.urandom(12)
        b = lambda x: base64.b64encode(x).decode()
        return json.dumps({"v": 1, "epk": b(epk), "iv": b(iv), "ct": b(AESGCM(key).encrypt(iv, json.dumps(obj).encode(), b"portfolio-mail-v1"))})
    DB[f"mail/{UID}"] = {"pkg": seal_mail({"uid": UID, "email": "friend@example.com", "refresh": "RT1", "pk8": pk8, "prefs": {"alerts": True, "weekly": True}})}
    # the engine's shared market data from the synthetic export
    eng = os.path.join(tmp, "engine"); sh = os.path.join(eng, "shared", "history"); os.makedirs(sh)
    unwrap = lambda p: (lambda x: x.get("data", x) if isinstance(x, dict) and isinstance(x.get("data"), dict) else x)(json.load(open(p)))
    json.dump(unwrap(os.path.join(SYN, "market", "latest.json")), open(os.path.join(eng, "shared", "latest.json"), "w"))
    json.dump(unwrap(os.path.join(SYN, "bench", "egx30.json")), open(os.path.join(eng, "shared", "bench.json"), "w"))
    for f in os.listdir(os.path.join(SYN, "history")):
        json.dump(unwrap(os.path.join(SYN, "history", f)), open(os.path.join(sh, f), "w"))
    # a second account that connected its Gmail: it signed up on 2026-08-15 holding 10 COMI and 9,145 cash (opening rows),
    # so August's statement (a deposit on the 1st, the COMI buy on the 2nd) is used from the 16th only
    gacct = ec.generate_private_key(ec.SECP256R1())
    gpk8 = base64.b64encode(gacct.private_bytes(serialization.Encoding.DER, serialization.PrivateFormat.PKCS8, serialization.NoEncryption())).decode()
    gpub = base64.b64encode(gacct.public_key().public_bytes(serialization.Encoding.X962, serialization.PublicFormat.UncompressedPoint)).decode()
    GUID = "Ugmail1"
    syn_settings = json.load(open(os.path.join(SYN, "portfolio", "settings.json")))
    syn_settings = syn_settings.get("data", syn_settings) if isinstance(syn_settings.get("data"), dict) else syn_settings
    gdocs = {
        "portfolio/settings": {**syn_settings, "name": "Friend Portfolio", "portfolioId": "acct-x", "inception": "2026-08", "trackFrom": "2026-08-15",
                               "cash": 9145, "cashDate": "2026-08-15", "account": {"holder": "Test Friend", "unifiedCode": ""}, "factsheetEmail": ""},   # private-scan: synthetic
        "portfolio/assets": {"items": {"Commercial International Bank": {"name": "Commercial International Bank", "symbol": "COMI", "sector": "Banks"}}},
        "portfolio/marks": {"months": {}},
        "ledger/y2026": {"rows": [
            {"id": "o1", "d": "2026-08-15", "t": "Deposit", "amt": 10000, "acc": "Main", "src": "manual", "opening": True},
            {"id": "o2", "d": "2026-08-15", "t": "Buy", "a": "Commercial International Bank", "q": 10, "p": 85.5, "amt": -855, "acc": "Main", "src": "manual", "opening": True}]},
        "sync/gmail": {"address": "friend.gmail@example.com", "appPassword": "abcdefghijklmnop"},
    }
    for k, v in gdocs.items():
        DB[f"users/{GUID}/docs/{k.replace('/', '__')}"] = {"blob": store.encode_doc({"pub": gpub}, k.split("/")[1], 1, v, "2026-08-15T10:00:00Z").decode().strip(), "updateTime": stamp()}
    DB[f"mail/{GUID}"] = {"pkg": seal_mail({"uid": GUID, "email": "friend2@example.com", "refresh": "RT2", "pk8": gpk8, "prefs": {"alerts": True, "weekly": False, "gmail": True}})}
    # the Thndr email: the synthetic statement PDFs in a DKIM-passed message from Thndr, as Gmail stores it
    pdfs = os.path.join(tmp, "pdfs")
    subprocess.run([sys.executable, os.path.join(ROOT, "tests", "fixtures", "make_statement_pdf.py"), pdfs, "--month", "2026-08"], check=True, capture_output=True)
    msg = EmailMessage()
    msg["Authentication-Results"] = "mx.google.com; dkim=pass header.i=@thndr.app header.s=s1 header.b=x; spf=pass smtp.mailfrom=system.thndr.app"
    msg["From"] = "Thndr <no-reply@system.thndr.app>"
    msg["To"] = "friend.gmail@example.com"
    msg["Subject"] = "Your monthly E-statement"
    msg.set_content("Your monthly statement is attached.")
    for f in ("account-statement.pdf", "position-snapshot.pdf"):
        msg.add_attachment(open(os.path.join(pdfs, f), "rb").read(), maintype="application", subtype="pdf", filename=f)
    RAW = base64.urlsafe_b64encode(msg.as_bytes()).decode()
    def statement_raw(month, *extra):
        d = os.path.join(tmp, "pdfs-" + month)
        subprocess.run([sys.executable, os.path.join(ROOT, "tests", "fixtures", "make_statement_pdf.py"), d, "--month", month, *extra], check=True, capture_output=True)
        m = EmailMessage()
        for k in ("Authentication-Results", "From", "To", "Subject"):
            m[k] = msg[k]
        m.set_content("Your monthly statement is attached.")
        for f in ("account-statement.pdf", "position-snapshot.pdf"):
            m.add_attachment(open(os.path.join(d, f), "rb").read(), maintype="application", subtype="pdf", filename=f)
        return base64.urlsafe_b64encode(m.as_bytes()).decode()
    RAW_SEP = statement_raw("2026-09", "--start", "9145", "--deposit", "1000", "--qty", "5", "--price", "90", "--close", "91", "--hold", "16")
    HIST_MAIL = {"h-aug": ("1788400000000", RAW), "h-sep": ("1791000000000", RAW_SEP)}   # private-scan: synthetic
    fetches = []
    def fake_fetch(after, seen, out_dir, addr=None, pw=None, query=None):
        fetches.append({"after": after, "addr": addr, "pw": pw, "query": "monthly" if query else "all"})
        if addr in ("hist@example.com", "wait@example.com"):
            os.makedirs(out_dir, exist_ok=True)
            ids = [] if addr == "wait@example.com" else ["h-aug", "h-sep"] if query else ["h-sep"]
            man = [{"id": i, "subject": "Your monthly E-statement", "date": HIST_MAIL[i][0]} for i in ids if i not in seen]
            for mm in man:
                json.dump({"id": mm["id"], "raw": HIST_MAIL[mm["id"]][1], "internalDate": mm["date"]}, open(os.path.join(out_dir, mm["id"] + ".json"), "w"))
            json.dump(man, open(os.path.join(out_dir, "manifest.json"), "w"))
            return {"found": len(man), "kept": len(man), "seen": 0, "otherSubject": 0}
        if pw != "abcdefghijklmnop":
            raise imap_fetch.FetchError("Gmail refused the app password (IMAP login)")
        os.makedirs(out_dir, exist_ok=True)
        man = [] if "18a0b0c0d0e0f001" in seen else [{"id": "18a0b0c0d0e0f001", "subject": "Your monthly E-statement", "date": "1788400000000"}]   # private-scan: synthetic
        if man:
            json.dump({"id": man[0]["id"], "raw": RAW, "internalDate": man[0]["date"]}, open(os.path.join(out_dir, man[0]["id"] + ".json"), "w"))
        json.dump(man, open(os.path.join(out_dir, "manifest.json"), "w"))
        return {"found": len(man), "kept": len(man), "seen": 0 if man else 1, "otherSubject": 0}
    imap_fetch.fetch = fake_fetch
    commits = []
    def gdoc(k):
        rec = DB.get(f"users/{GUID}/docs/{k.replace('/', '__')}")
        return json.loads(store.unseal(gacct, rec["blob"]).decode())["data"] if rec else None
    # two accounts made with "Build it from my Thndr emails": one whose Gmail has monthly statements, one with none yet
    def new_account(uid, email, gmail, refresh):
        k = ec.generate_private_key(ec.SECP256R1())
        pub = base64.b64encode(k.public_key().public_bytes(serialization.Encoding.X962, serialization.PublicFormat.UncompressedPoint)).decode()
        pk = base64.b64encode(k.private_bytes(serialization.Encoding.DER, serialization.PrivateFormat.PKCS8, serialization.NoEncryption())).decode()
        docs = {"portfolio/settings": {**syn_settings, "name": "History Portfolio", "portfolioId": "acct-h", "inception": "2026-09", "cash": 0, "cashDate": "2026-09-20",
                                       "account": {"holder": "Test Friend", "unifiedCode": ""}, "factsheetEmail": "", "historyImport": {"status": "pending"}},   # private-scan: synthetic
                "portfolio/assets": {"items": {}}, "portfolio/marks": {"months": {}}, "ledger/y2026": {"rows": []},
                "sync/gmail": {"address": gmail, "appPassword": "historyhistoryhi"}}
        for kk, v in docs.items():
            DB[f"users/{uid}/docs/{kk.replace('/', '__')}"] = {"blob": store.encode_doc({"pub": pub}, kk.split("/")[1], 1, v, "2026-09-20T10:00:00Z").decode().strip(), "updateTime": stamp()}
        DB[f"mail/{uid}"] = {"pkg": seal_mail({"uid": uid, "email": email, "refresh": refresh, "pk8": pk, "prefs": {"alerts": True, "weekly": False, "gmail": True}})}
        return k
    HUID, WUID = "Uhistory", "Uwaiting"
    hacct = new_account(HUID, "hist2@example.com", "hist@example.com", "RT4")
    wacct = new_account(WUID, "wait2@example.com", "wait@example.com", "RT5")
    def adoc(uid, key, k):
        rec = DB.get(f"users/{uid}/docs/{k.replace('/', '__')}")
        return json.loads(store.unseal(key, rec["blob"]).decode())["data"] if rec else None
    # the site owner's own account: shares the MAIN portfolio (the engine's documents, from the synthetic export)
    oacct = ec.generate_private_key(ec.SECP256R1())
    opk8 = base64.b64encode(oacct.private_bytes(serialization.Encoding.DER, serialization.PrivateFormat.PKCS8, serialization.NoEncryption())).decode()
    OUID = "Uowner"
    store.migrate(SYN, eng, os.path.join(code, "p", "khaled", "keys.json"))
    DB[f"users/{OUID}/docs/portfolio__settings"] = {"blob": store.encode_doc({"pub": base64.b64encode(oacct.public_key().public_bytes(serialization.Encoding.X962, serialization.PublicFormat.UncompressedPoint)).decode()}, "settings", 1, {"name": "Owner account"}, "2026-09-20T00:00:00Z").decode().strip(), "updateTime": stamp()}
    ram.OWNER_HASH = hashlib.sha256(b"owner@example.com").hexdigest()     # a stand-in owner
    DB[f"mail/{OUID}"] = {"pkg": seal_mail({"uid": OUID, "email": "owner@example.com", "refresh": "RT3", "pk8": opk8, "prefs": {"alerts": True, "weekly": True, "gmail": True, "shareMain": True}})}
    # account 1 asks for the main portfolio too: it is not the owner, so it must share its own
    DB[f"mail/{UID}"] = {"pkg": seal_mail({"uid": UID, "email": "friend@example.com", "refresh": "RT1", "pk8": pk8, "prefs": {"alerts": True, "weekly": True, "shareMain": True}})}
    CLAIMS = {UID: {"email": "friend@example.com", "email_verified": True}, GUID: {"email": "friend.gmail@example.com", "email_verified": True},
              OUID: {"email": "owner@example.com", "email_verified": True}}
    def jwt(u):
        return "h." + base64.urlsafe_b64encode(json.dumps({"user_id": u, **CLAIMS.get(u, {})}).encode()).decode().rstrip("=") + ".s"
    def uid_of(tok):
        try:
            part = tok.split(".")[1]
            return json.loads(base64.urlsafe_b64decode(part + "=" * (-len(part) % 4)))["user_id"]
        except Exception:
            return None
    gpub_ = gpub
    opub_ = base64.b64encode(oacct.public_key().public_bytes(serialization.Encoding.X962, serialization.PublicFormat.UncompressedPoint)).decode()
    LINKS = {(UID, "Uzeyad"): {"status": "received", "name": "Zeyad's Portfolio", "pub": gpub_, "email": "zeyad@example.com"},
             (UID, GUID): {"status": "friends", "name": "Friend Portfolio", "pub": gpub_, "email": "friend2@example.com"},
             (GUID, UID): {"status": "friends", "name": "Demo", "pub": apub, "email": "friend@example.com"},
             (OUID, GUID): {"status": "friends", "name": "Friend Portfolio", "pub": gpub_, "email": "friend2@example.com"},
             (GUID, OUID): {"status": "friends", "name": "Owner", "pub": opub_, "email": "owner@example.com"},
             # someone posing as the friend account (their own key, the friend's email): must never get a copy
             (UID, "Uimposter"): {"status": "friends", "name": "Friend Portfolio", "pub": opub_, "email": "friend2@example.com"}}
    # directory/{email}: only the owner of that sign-in email writes it (the database rules), so it is the truth
    DIRECTORY = {"friend@example.com": {"uid": UID, "pub": apub}, "friend2@example.com": {"uid": GUID, "pub": gpub_}, "owner@example.com": {"uid": OUID, "pub": opub_}}
    SHARES, STATUS, share_writes = {}, {}, []
    PROFILES, profile_writes = {OUID: {"name": {"stringValue": "Owner"}}, UID: {"name": {"stringValue": "Demo"}}}, []
    def open_share(priv, key):
        e = json.loads(SHARES[key]["pkg"]); b = base64.b64decode; epk = b(e["epk"])
        k = HKDF(hashes.SHA256(), 32, epk, b"portfolio-share-v1").derive(priv.exchange(ec.ECDH(), ec.EllipticCurvePublicKey.from_encoded_point(ec.SECP256R1(), epk)))
        return json.loads(gzip.decompress(AESGCM(k).decrypt(b(e["iv"]), b(e["ct"]), b"portfolio-share-v1")))
    FS = ram.FS
    calls = []
    class FakeHttp:
        def json(self, method, url, body=None, headers=None, form=False):
            calls.append((method, url.split("?")[0].replace(FS, "")))
            tok = (headers or {}).get("Authorization", "").replace("Bearer ", "")
            if url.startswith("https://securetoken"):
                u = {"RT1": UID, "RT2": GUID, "RT3": OUID, "RT4": HUID, "RT5": WUID}.get(body.get("refresh_token"))
                return (200, {"id_token": jwt(u), "user_id": u}) if u else (400, {"error": {"message": "INVALID_REFRESH_TOKEN"}})
            if url.startswith(FS + ":commit"):
                commits.append((uid_of(tok), len(body["writes"])))
                for w in body["writes"]:
                    k = w["update"]["name"].split("/documents/", 1)[1]
                    if uid_of(tok) != k.split("/")[1]: return 403, {"error": {"status": "PERMISSION_DENIED"}}
                    cur, pre = DB.get(k), w["currentDocument"]
                    if ("exists" in pre and bool(cur) != pre["exists"]) or ("updateTime" in pre and (not cur or cur["updateTime"] != pre["updateTime"])):
                        return 400, {"error": {"status": "FAILED_PRECONDITION"}}
                for w in body["writes"]:
                    DB[w["update"]["name"].split("/documents/", 1)[1]] = {"blob": w["update"]["fields"]["blob"]["stringValue"], "updateTime": stamp()}
                return 200, {"writeResults": []}
            path = url.split("?")[0].replace(FS + "/", "")
            if path == "mail":
                return 200, {"documents": [{"name": f"projects/p/databases/(default)/documents/mail/{k.split('/')[1]}", "fields": {"pkg": {"stringValue": v["pkg"]}}} for k, v in DB.items() if k.startswith("mail/")]}
            if path.startswith("users/") and path.count("/") == 1:      # the profile document (users/{uid}): only its account
                u = path.split("/")[1]
                if uid_of(tok) != u: return 403, {"error": {"status": "PERMISSION_DENIED"}}
                if method == "GET":
                    return (200, {"fields": PROFILES[u]}) if u in PROFILES else (404, {"error": {"status": "NOT_FOUND"}})
                for fp in urllib.parse.parse_qs(url.split("?", 1)[1] if "?" in url else "").get("updateMask.fieldPaths", []):
                    if fp not in body["fields"]:
                        PROFILES.setdefault(u, {}).pop(fp, None)      # a masked field missing from the body is deleted
                PROFILES.setdefault(u, {}).update(body["fields"]); profile_writes.append(u)
                return 200, {}
            if path.startswith("users/"):
                uid = path.split("/")[1]
                if uid_of(tok) != uid: return 403, {"error": {"status": "PERMISSION_DENIED"}}
                if method == "GET":
                    return 200, {"documents": [{"name": "x/" + k, "fields": {"blob": {"stringValue": v["blob"]}}, "updateTime": v["updateTime"]} for k, v in sorted(DB.items()) if k.startswith(path + "/")]}
                if method == "PATCH":
                    q = url.split("?", 1)[1]
                    cur = DB.get(path)
                    if ("exists=false" in q and cur) or ("updateTime=" in q and (not cur or cur["updateTime"] not in q.replace("%3A", ":"))):
                        return 400, {"error": {"status": "FAILED_PRECONDITION"}}
                    DB[path] = {"blob": body["fields"]["blob"]["stringValue"], "updateTime": stamp()}
                    return 200, {"updateTime": DB[path]["updateTime"]}
            deny = (403, {"error": {"status": "PERMISSION_DENIED"}})
            if path == "status":
                if uid_of(tok) != OUID: return deny
                return 200, {"documents": [{"name": f"x/status/{u}", "fields": {k: {"stringValue": v} for k, v in f.items()}} for u, f in STATUS.items()]}
            if path.startswith("status/"):
                u = path.split("/")[1]
                if uid_of(tok) != u or method != "PATCH": return deny
                STATUS.setdefault(u, {}).update({k: v["stringValue"] for k, v in body["fields"].items()})
                return 200, {}
            if path.startswith("directory/"):
                d = DIRECTORY.get(urllib.parse.unquote(path.split("/", 1)[1]))
                if not tok or method != "GET": return deny
                return (200, {"fields": {k: {"stringValue": v} for k, v in d.items()}}) if d else (404, {"error": {"status": "NOT_FOUND"}})
            if path.startswith("links/"):
                u = path.split("/")[1]
                if uid_of(tok) != u: return deny
                return 200, {"documents": [{"name": f"x/links/{a}/with/{b}", "fields": {k: {"stringValue": v} for k, v in f.items()}} for (a, b), f in sorted(LINKS.items()) if a == u]}
            if path.startswith("shares/"):
                _, owner, _, viewer = path.split("/")
                if method == "GET":
                    if uid_of(tok) not in (owner, viewer): return deny
                    return (200, {"fields": {k: {"stringValue": v} for k, v in SHARES[(owner, viewer)].items()}}) if (owner, viewer) in SHARES else (404, {})
                if uid_of(tok) != owner or (LINKS.get((owner, viewer)) or {}).get("status") != "friends": return deny
                SHARES[(owner, viewer)] = {k: v["stringValue"] for k, v in body["fields"].items()}
                share_writes.append((owner, viewer))
                return 200, {}
            return 404, {}
    sent = []
    def send(to, subj, text, html, att=None):
        sent.append({"to": to, "subject": subj, "text": text, "html": html or "", "att": att or []})
        if os.environ.get("DUMP_EMAILS"):      # every email this test sends, to read them as a person would
            d = os.environ["DUMP_EMAILS"]; os.makedirs(d, exist_ok=True); n = len([f for f in os.listdir(d) if f.endswith('.txt')]) + 1
            open(os.path.join(d, f"{n:02d}.txt"), "w").write(f"To: {to}\nSubject: {subj}\nAttachments: {[a[0] for a in (att or [])]}\n\n{text}")
            open(os.path.join(d, f"{n:02d}.html"), "w").write(html or "")
            for a in att or []:
                open(os.path.join(d, f"{n:02d}-{a[0]}"), "wb").write(a[1])
    # the site repository the month-end files are published to (a local bare repository standing in for GitHub)
    site_bare, site_seed = os.path.join(tmp, "site.git"), os.path.join(tmp, "site-seed")
    subprocess.run(["git", "init", "-q", "--bare", "-b", "main", site_bare], check=True)
    subprocess.run(["git", "clone", "-q", site_bare, site_seed], check=True, capture_output=True)
    open(os.path.join(site_seed, "README.md"), "w").write("site\n")
    for c in (["add", "-A"], ["-c", "user.name=t", "-c", "user.email=t@example.com", "commit", "-q", "-m", "seed"], ["push", "-q", "origin", "HEAD:main"]):
        subprocess.run(["git", "-C", site_seed] + c, check=True, capture_output=True)
    os.environ["SITE_REMOTE"] = site_bare
    argv = ["--engine", eng, "--code", code, "--now", "2026-09-24T19:30:00Z"]     # Thursday 22:30 Cairo
    rc = ram.main(argv, http=FakeHttp(), send=send)
    check("the job succeeds", rc == 0)
    gsent = [m for m in sent if m["to"] == "friend2@example.com"]
    hsent = [m for m in sent if m["to"] in ("hist2@example.com", "wait2@example.com")]
    check("the owner account (Gmail and emails switched on) gets no import and no email: it only shares",
          not [m for m in sent if m["to"] == "owner@example.com"] and all(f["addr"] in ("friend.gmail@example.com", "hist@example.com", "wait@example.com") for f in fetches), json.dumps([m["subject"] for m in sent if m["to"] == "owner@example.com"]))
    sent[:] = [m for m in sent if m["to"] == "friend@example.com"]
    check("account 1: new heads-up items, the weekly summary, then the friend request, to its own address",
          len(sent) == 3 and "heads-up" in sent[0]["subject"] and sent[1]["html"] and sent[2]["subject"] == "Zeyad's Portfolio wants to be friends on the portfolio site"
          and "?friends" in sent[2]["text"], json.dumps([m["subject"] for m in sent]))
    sent[:] = sent[:2]
    # one login: the owner's verified account now holds the MAIN portfolio's key, sealed to the account's key
    mk = (PROFILES.get(OUID) or {}).get("mainKey", {}).get("stringValue")
    opened_mk = ram.open_json(oacct, mk, b"portfolio-mainkey-v1") if mk else {}
    main_pk8 = store.unlock(store.load_keys(os.path.join(code, "p", "khaled", "keys.json")), os.environ["SETUP_KEY"]).private_bytes(
        serialization.Encoding.DER, serialization.PrivateFormat.PKCS8, serialization.NoEncryption())
    check("one login: the owner's account gets the main portfolio's key, sealed to its own key; no other account does",
          opened_mk.get("id") == "khaled" and base64.b64decode(opened_mk.get("pk8", "")) == main_pk8 and profile_writes == [OUID]
          and "mainKey" not in (PROFILES.get(UID) or {}), json.dumps(profile_writes))
    # friends
    check("friends: a copy is written for every friend and nobody else (not for a link posing as a friend with another key)", sorted(share_writes) == sorted([(UID, GUID), (GUID, UID), (GUID, OUID), (OUID, GUID)]), json.dumps(share_writes))
    s1 = open_share(acct, (GUID, UID))
    p1 = s1.get("profile") or {}
    check("friends: the Gmail account's copy opens with its friend's key and is its percentages profile (months, holdings by weight, its trades)",
          s1.get("v") == 2 and s1["name"] == "Friend Portfolio" and p1.get("months") and [h["s"] for h in p1.get("holdings") or []] == ["COMI"]
          and [(t["side"], t["kind"], t["s"]) for t in p1.get("trades") or []] == [("buy", "new", "COMI")], json.dumps(s1)[:300])
    # percentages only: no amount, share count, price or document anywhere in it
    j1s = json.dumps(s1)
    check("friends: the copy holds percentages only (no amounts, share counts, prices, documents, Thndr account or Gmail login)",
          not any(k in j1s for k in ('"amt"', '"q"', '"p"', '"mv"', '"cash"', '"docs"', '"account"', "abcdefghijklmnop")) and all(abs(h["w"]) <= 1 for h in p1["holdings"]), j1s[:300])
    s2 = open_share(gacct, (UID, GUID))
    check("friends: an account that is not the owner asking for the main portfolio shares its own profile", s2.get("v") == 2 and s2["name"] != "Main" and (s2.get("profile") or {}).get("months"))
    s3 = open_share(gacct, (OUID, GUID))
    check("friends: the owner's verified account shares the MAIN portfolio's profile",
          s3.get("v") == 2 and (s3.get("profile") or {}).get("months") and s3["name"] != "Owner account", json.dumps(s3)[:200])
    jg, j1 = json.loads(STATUS.get(GUID, {}).get("job", "{}")), json.loads(STATUS.get(UID, {}).get("job", "{}"))
    check("admin status: the job writes each account's line (Gmail result, month-end report, friends), no figures (account 1 lists 2 friend links: its friend and the impostor link)",
          jg.get("gmail", {}).get("ok") and jg.get("friends") == 2 and jg.get("report") == "Aug-26 sent 2026-09-24" and j1.get("friends") == 2 and not j1.get("gmail"), json.dumps([jg, j1]))
    # the Gmail account
    check("gmail: its own Gmail login is used, searching from 3 days before tracking started", fetches and fetches[0] == {"after": "2026/08/12", "addr": "friend.gmail@example.com", "pw": "abcdefghijklmnop", "query": "all"}, json.dumps(fetches))
    rows = (gdoc("ledger/y2026") or {}).get("rows") or []
    check("gmail: the opening rows are kept and the statement rows before the tracking start are not added again",
          sorted(r["id"] for r in rows) == ["o1", "o2"], json.dumps(rows)[:400])
    mk = ((gdoc("portfolio/marks") or {}).get("months") or {}).get("2026-08") or {}
    check("gmail: August is posted from the statement (month-end cash and securities, final)", mk.get("source") == "statement" and not mk.get("provisional") and abs(mk.get("cash", 0) - 9145) < 0.01 and abs(mk.get("securities", 0) - 860) < 0.01, json.dumps(mk))
    check("gmail: the Thndr account code is recorded from the statement", ((gdoc("portfolio/settings") or {}).get("account") or {}).get("unifiedCode") == "1234567")   # private-scan: synthetic
    check("gmail: the import and sync state are saved in ONE commit (the month-end stamp is a second)", [n for u, n in commits if u == GUID][1:] == [1] and len([n for u, n in commits if u == GUID]) == 2 and (gdoc("imports/2026-08") or {}).get("fullMonth") and "18a0b0c0d0e0f001" in ((gdoc("sync/state") or {}).get("seen") or {}), json.dumps(commits))
    check("gmail: the summary email goes to the account's own address", len(gsent) == 2 and "Aug-26 statement posted" in gsent[0]["subject"], json.dumps([m["subject"] for m in gsent]))
    rep = gsent[1] if len(gsent) > 1 else {"subject": "", "att": []}
    names = [a[0] for a in rep["att"]]
    check("month-end: the Aug-26 report goes to the account's own address with the Excel workbook attached",
          rep["subject"] == "Friend Portfolio · month-end report Aug-26" and "FriendPortfolio-Aug-26.xlsx" in names, json.dumps([rep["subject"], names]))
    if names:
        import io
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(next(a for a in rep["att"] if a[0].endswith(".xlsx"))[1]), data_only=True)
        check("month-end: the workbook opens and starts with the Summary sheet", wb.sheetnames[0] == "Summary", str(wb.sheetnames))
    if os.environ.get("KEEP_REPORT"):
        for a in rep["att"]:
            open(os.path.join(os.environ["KEEP_REPORT"], a[0]), "wb").write(a[1])
    has_pw = subprocess.run(["node", "-e", "require.resolve('playwright', {paths: [process.argv[1]]})", os.path.join(PDFJS_TOOLS, "node_modules")], capture_output=True).returncode == 0
    if has_pw or os.environ.get("REQUIRE_PDF") == "1":
        check("month-end: the PDF factsheet is attached and the email shows the headline figures (value, the month's return, top holding)",
              "FriendPortfolio-Aug-26.pdf" in names and rep["html"] and next(a for a in rep["att"] if a[0].endswith(".pdf"))[1][:5] == b"%PDF-"
              and "Value at month-end" in rep["text"] and "10,005 EGP" in rep["text"] and "Return in Aug" in rep["text"] and "COMI" in rep["text"], json.dumps(names) + rep["text"][:600])   # private-scan: synthetic
    else:
        print("SKIP month-end PDF: no Playwright next to the tools (CI step 5f checks it)")
    # the same files as downloads on the site: sealed to the account's key under a/<hash>/exports/, the list sealed too
    chk = os.path.join(tmp, "site-check")
    subprocess.run(["git", "clone", "-q", site_bare, chk], check=True, capture_output=True)
    fdir = os.path.join(chk, ram.files_dir(GUID))
    idx = json.loads(store.unseal(gacct, open(os.path.join(fdir, "index.enc.json"), "rb").read())) if os.path.exists(os.path.join(fdir, "index.enc.json")) else []
    got = store.unseal(gacct, open(os.path.join(chk, idx[0]["file"]), "rb").read()) if idx else b""
    check("month-end files on the site: the Aug-26 workbook sealed to the account's own key, listed in a sealed index, nothing readable in the clear",
          idx and idx[0]["month"] == "2026-08" and idx[0]["name"] == "FriendPortfolio-Aug-26.xlsx" and got == next(a for a in rep["att"] if a[0].endswith(".xlsx"))[1]
          and not [f for f in os.listdir(fdir) if not f.endswith(".enc.json")] and GUID not in ram.files_dir(GUID), json.dumps(idx))
    imp = gdoc("imports/2026-08") or {}
    check("month-end: the month is stamped as sent in the account", (imp.get("reports") or {}).get("emailedAt") and "reportsPending" not in imp, json.dumps(imp)[:300])
    # history import
    hs = [m for m in hsent if m["to"] == "hist2@example.com"]
    hrows = (adoc(HUID, hacct, "ledger/y2026") or {}).get("rows") or []
    hset = adoc(HUID, hacct, "portfolio/settings") or {}
    hmk = (adoc(HUID, hacct, "portfolio/marks") or {}).get("months") or {}
    check("history: the monthly statements are looked for first, then everything after the starting point",
          [f["query"] for f in fetches if f["addr"] == "hist@example.com"] == ["monthly", "all"] and [f["after"] for f in fetches if f["addr"] == "hist@example.com"] == ["2019/01/01", "2026/09/30"], json.dumps([f for f in fetches if f["addr"] == "hist@example.com"]))
    check("history: the earliest statement (Aug-26) is the starting point: its holding and cash as opening rows on its last day",
          [(r["d"], r["t"], r.get("q"), r.get("opening")) for r in hrows if r.get("opening")] == [("2026-08-31", "Deposit", None, True), ("2026-08-31", "Buy", 10, True)]
          and abs([r for r in hrows if r.get("opening") and r["t"] == "Deposit"][0]["amt"] - 10005) < 0.01, json.dumps(hrows)[:400])
    check("history: the later statement (Sep-26) is built on top: its deposit and buy, then the snapshot's extra share and the cash as adjustments",
          sorted((r["d"], r["t"], r.get("src")) for r in hrows if not r.get("opening")) == [("2026-09-01", "Deposit", "history-2026-09"), ("2026-09-02", "Buy", "history-2026-09"), ("2026-09-30", "Buy", "history-adjust"), ("2026-09-30", "Deposit", "history-adjust")]
          and [(r.get("q"), r.get("amt")) for r in hrows if r.get("src") == "history-adjust"] == [(1, -91), (None, 91)] and all(r["note"].startswith("Adjustment") for r in hrows if r.get("src") == "history-adjust"), json.dumps(hrows)[:600])
    check("history: settings start at Aug-26 with the account code; the import is marked done",
          hset.get("inception") == "2026-08" and hset.get("trackFrom") == "2026-09-30" and (hset.get("historyImport") or {}).get("status") == "done" and (hset.get("historyImport") or {}).get("adjustments") == 2 and (hset.get("account") or {}).get("unifiedCode") == "1234567", json.dumps({k: hset.get(k) for k in ("inception", "trackFrom", "historyImport", "account")}))   # private-scan: synthetic
    check("history: both months' marks come from the statements", (hmk.get("2026-08") or {}).get("source") == "statement" and abs((hmk.get("2026-09") or {}).get("cash", 0) - 9695) < 0.01 and abs((hmk.get("2026-09") or {}).get("securities", 0) - 1456) < 0.01, json.dumps(hmk))
    check("history: one summary email, and only the latest month's report (Sep-26), not the old one",
          [m["subject"] for m in hs] == ["History Portfolio: built from your Thndr emails", "History Portfolio · month-end report Sep-26"] and "Starting point: Aug-26 statement" in hs[0]["text"] and "Monthly statements used: 2 (up to Sep-26)" in hs[0]["text"] and "Adjustments: 2 (Sep-26)" in hs[0]["text"]
          and ((adoc(HUID, hacct, "imports/2026-08") or {}).get("reports") or {}).get("emailedAt") == "skipped (history import)", json.dumps([m["subject"] for m in hs]))
    ws = [m for m in hsent if m["to"] == "wait2@example.com"]
    wstate = adoc(WUID, wacct, "sync/mail") or {}
    check("history: with no monthly statement yet, the account is told once and its portfolio is untouched",
          [m["subject"] for m in ws] == ["History Portfolio: waiting for a monthly Thndr statement"] and (wstate.get("history") or {}).get("status") == "waiting"
          and (adoc(WUID, wacct, "portfolio/settings") or {}).get("historyImport") == {"status": "pending"} and not [u for u, n in commits if u == WUID], json.dumps([m["subject"] for m in ws]))
    gm = (gdoc("sync/mail") or {}).get("gmail") or {}
    check("gmail: the result is recorded for the site (ok, 1 new, 1 applied)", gm.get("ok") and gm.get("new") == 1 and gm.get("applied") == 1, json.dumps(gm))
    check("the heads-up email lists the synthetic ex-dividend and target items", "Ex-dividend on " in sent[0]["text"] and "reached its target" in sent[0]["text"]
          and "Ex-dividend: Ex-dividend" not in sent[0]["text"] and "Target reached:" not in sent[0]["text"], sent[0]["text"])
    rec = DB.get(f"users/{UID}/docs/sync__mail")
    st = json.loads(store.unseal(acct, rec["blob"]).decode())["data"] if rec else {}
    check("what was sent is saved back to the account, encrypted to its key", len(st.get("alertsSent", {})) == 2 and st.get("weeklySent") == "2026-09-24")
    check("the account that did not opt in is never read", not any(u.startswith(f"/users/{OTHER}") for _, u in calls))
    sent.clear(); commits.clear(); share_writes.clear()
    rc2 = ram.main(argv, http=FakeHttp(), send=send)
    check("friends: a second run rewrites no unchanged copy and emails no request again", share_writes == [], json.dumps(share_writes))
    check("one login: written once", profile_writes == [OUID], json.dumps(profile_writes))
    check("a second run sends nothing (each alert once, one summary a week)", rc2 == 0 and sent == [], json.dumps([m["subject"] for m in sent]))
    check("gmail: a second run finds nothing new and changes no portfolio document", [n for u, n in commits if u == GUID] == [1] and sorted(r["id"] for r in (gdoc("ledger/y2026") or {}).get("rows") or []) == ["o1", "o2"], json.dumps(commits))
    STATUS["Unewbie"] = {"name": "Newbie's Portfolio", "email": "newbie@example.com"}
    sent.clear(); ram.main(argv, http=FakeHttp(), send=send)
    ns = [m for m in sent if m["to"] == "owner@example.com"]
    check("the owner is emailed once about a new account (name and email)", len(ns) == 1 and ns[0]["subject"] == "New on your portfolio site: Newbie's Portfolio" and "newbie@example.com" in ns[0]["text"], json.dumps([m["subject"] for m in sent]))
    sent.clear(); ram.main(argv, http=FakeHttp(), send=send)
    check("... and not again", not [m for m in sent if m["to"] == "owner@example.com"])
    # a wrong app password: emailed once, recorded, the portfolio untouched
    before = {k: v["updateTime"] for k, v in DB.items() if k.startswith(f"users/{GUID}/docs/") and not k.endswith("sync__mail")}
    login = DB[f"users/{GUID}/docs/sync__gmail"]
    DB[f"users/{GUID}/docs/sync__gmail"] = {"blob": store.encode_doc({"pub": gpub}, "gmail", 2, {"address": "friend.gmail@example.com", "appPassword": "wrongwrongwrongw"}, "2026-09-24T00:00:00Z").decode().strip(), "updateTime": stamp()}
    before[f"users/{GUID}/docs/sync__gmail"] = DB[f"users/{GUID}/docs/sync__gmail"]["updateTime"]
    for _ in range(2):
        sent.clear(); ram.main(argv, http=FakeHttp(), send=send)
        if _ == 0:
            first_err = [m["subject"] for m in sent if m["to"] == "friend2@example.com"]
    gm = (gdoc("sync/mail") or {}).get("gmail") or {}
    check("gmail: a refused app password is emailed once and recorded for the site",
          first_err == ["Friend Portfolio: Thndr emails could not be read"] and not [m for m in sent if m["to"] == "friend2@example.com"] and gm.get("ok") is False and "refused" in gm.get("error", ""), json.dumps([first_err, gm]))
    check("gmail: nothing in the portfolio changed", before == {k: v["updateTime"] for k, v in DB.items() if k.startswith(f"users/{GUID}/docs/") and not k.endswith("sync__mail")})
    DB[f"users/{GUID}/docs/sync__gmail"] = login
    # a step failing after an email went out: what was sent is still recorded, so the next run does not send it again
    DB.pop(f"users/{UID}/docs/sync__mail", None)     # account 1 starts over: its heads-up items are new again
    real_run = ram.subprocess.run
    def failing_weekly(cmd, *a, **k):
        if any(str(x).endswith("weekly.js") for x in cmd):
            return subprocess.CompletedProcess(cmd, 1, "", "weekly.js broke")
        return real_run(cmd, *a, **k)
    ram.subprocess.run = failing_weekly
    try:
        sent.clear(); ram.main(argv, http=FakeHttp(), send=send); first = [m["subject"] for m in sent if m["to"] == "friend@example.com"]
        sent.clear(); ram.main(argv, http=FakeHttp(), send=send); again = [m["subject"] for m in sent if m["to"] == "friend@example.com"]
    finally:
        ram.subprocess.run = real_run
    check("a step failing after the heads-up email went out (the weekly here): the email is recorded, the next run does not repeat it",
          any("heads-up" in x for x in first) and not any("heads-up" in x for x in again), json.dumps([first, again]))
    upkg = DB[f"mail/{UID}"]["pkg"]
    DB[f"mail/{UID}"]["pkg"] = seal_mail({"uid": "Usomeoneelse", "email": "x@example.com", "refresh": "RT1", "pk8": pk8})
    sent.clear()
    ram.main(argv, http=FakeHttp(), send=send)
    check("a package naming another account is refused", sent == [])
    DB[f"mail/{UID}"]["pkg"] = upkg

    # ---- the monthly friends leaderboard: in the first days of a month, to each account's own address, percentages only ----
    def lb_run(now):
        sent.clear(); ram.main(["--engine", eng, "--code", code, "--now", now], http=FakeHttp(), send=send)
        return {m["to"]: m for m in sent if "leaderboard" in m["subject"]}
    lb = lb_run("2026-09-03T13:30:00Z")
    rc3 = {m["to"]: m for m in sent if " report card" in m["subject"]}
    g = lb.get("friend2@example.com") or {}
    check("leaderboard: on the 3rd, an account whose friends' copies all cover August gets it (the owner too, ranked on the main portfolio)",
          g.get("subject", "").startswith("August 2026 leaderboard: you are #") and g["subject"].endswith(" of 3") and "Demo" in g["text"] and "Owner" in g["text"]
          and (lb.get("owner@example.com") or {}).get("subject", "").endswith(" of 2"), json.dumps({k: v["subject"] for k, v in lb.items()}))
    check("leaderboard: an account with a friend who has no copy for it waits (until the 8th)", "friend@example.com" not in lb)
    aug = {str(k): ram.month_figures(open_share({UID: acct, GUID: gacct, OUID: oacct}[k[1]], k)["profile"], "2026-08") for k in [(UID, GUID), (OUID, GUID)]}
    pct = lambda x: ("+" if x > 0.00005 else "−" if x < -0.00005 else "") + f"{abs(x) * 100:.1f}%"
    check("leaderboard: each friend's August return and year so far come from the profile they share",
          all(v[0] is not None and pct(v[0]) in g.get("text", "") and pct(v[1]) in g.get("text", "") for v in aug.values()), json.dumps(aug))
    check("leaderboard: percentages only (no EGP, no amounts)", lb and all("EGP" not in m["text"] and "EGP" not in m["html"] for m in lb.values()))
    check("leaderboard: not sent twice", lb_run("2026-09-03T17:30:00Z") == {})
    check("report card: not sent twice", not [m for m in sent if " report card" in m["subject"] and m["to"] == "friend2@example.com"])
    lb = lb_run("2026-09-08T13:30:00Z")
    rc8 = {m["to"]: m for m in sent if " report card" in m["subject"]}
    check("report card: on the 3rd, only the account whose August statement is in gets it (to its own address); the others on the 5th or later",
          sorted(rc3) == ["friend2@example.com", "hist2@example.com"] and rc3["friend2@example.com"]["subject"].startswith("Friend Portfolio: August 2026 report card · ")
          and sorted(rc8) == ["friend@example.com"], json.dumps([list(rc3), list(rc8)]))
    check("report card: an account with nothing in the month (still waiting to be built) gets none, and it is not asked again",
          (adoc(WUID, wacct, "sync/mail") or {}).get("reportCardSent") == "2026-08")
    check("report card: the site owner's sign-in account (no portfolio of its own) never gets one, the history accounts' empty months neither break the run",
          "owner@example.com" not in rc3 and "owner@example.com" not in rc8)
    check("report card: the account's own figures (its sales and its month), recorded as sent",
          "AUGUST 2026 NEXT TO JULY 2026" in rc8["friend@example.com"]["text"] and "Sales  0  1" in rc8["friend@example.com"]["text"] and (adoc(UID, acct, "sync/mail") or {}).get("reportCardSent") == "2026-08", rc8["friend@example.com"]["text"][:400])
    check("leaderboard: on the 8th the waiting account gets it, with the friends whose copies are there", list(lb) == ["friend@example.com"]
          and lb["friend@example.com"]["subject"].endswith(" of 2") and "Friend Portfolio" in lb["friend@example.com"]["text"], json.dumps({k: v["subject"] for k, v in lb.items()}))
    st1 = adoc(UID, acct, "sync/mail") or {}
    check("leaderboard: the month is recorded for the account", st1.get("leaderboardSent") == "2026-08", json.dumps(st1.get("leaderboardSent")))
    # the ranking itself
    me = {"months": [{"m": "2026-08", "r": 0.02, "b": 0.01}, {"m": "2026-07", "r": 0.1, "b": 0}], "trades": [{"d": "2026-08-10", "side": "sell", "s": "COMI", "ret": 0.12}], "asOf": "2026-09-01"}
    a = {"months": [{"m": "2026-08", "r": 0.05, "b": 0.01}], "trades": [{"d": "2026-08-20", "side": "sell", "s": "SWDY", "ret": 0.3}, {"d": "2026-07-20", "side": "sell", "s": "X", "ret": 0.9}], "asOf": "2026-08-31"}
    old = {"months": [{"m": "2026-08", "r": 0.5, "b": 0.01}], "trades": [], "asOf": "2026-08-15"}
    rows, bench, best = ram.leaderboard(me, [("Omar", a), ("Old copy", old), ("Nobody", {"months": []})], "2026-08")
    check("leaderboard: ranked by the month's return, a copy from mid-month and one without the month last, the index from the account",
          [r["who"] for r in rows] == ["Omar", "You", "Old copy", "Nobody"] and rows[2]["m"] is None and abs(rows[1]["ytd"] - (1.1 * 1.02 - 1)) < 1e-9 and bench == 0.01, json.dumps(rows))
    check("leaderboard: a copy counts for September only from its last session (Wed 30 Sep), for October from Thu 29 Oct (the 30th and 31st are Fri-Sat)",
          ram.fresh_for({"asOf": "2026-09-30T15:10:00+03:00"}, "2026-09") and not ram.fresh_for({"asOf": "2026-09-29T15:10:00+03:00"}, "2026-09")
          and ram.fresh_for({"asOf": "2026-10-29"}, "2026-10") and not ram.fresh_for({"asOf": "2026-10-28"}, "2026-10"))
    check("leaderboard: the best sale of the month among the ranked (not another month's)", best == {"who": "Omar", "s": "SWDY", "ret": 0.3}, json.dumps(best))
    gk = ec.generate_private_key(ec.SECP256R1())
    gpk = base64.b64encode(gk.public_key().public_bytes(serialization.Encoding.X962, serialization.PublicFormat.UncompressedPoint)).decode()
    check("open_json opens what seal_json sealed", ram.open_json(gk, ram.seal_json({"v": 2, "x": [1]}, gpk, b"portfolio-share-v1"), b"portfolio-share-v1") == {"v": 2, "x": [1]})

    # ---- the yearly wrap-up: early January, last year, ranked among friends; the owner's from the main portfolio ----
    def wrap_run(now):
        sent.clear(); ram.main(["--engine", eng, "--code", code, "--now", now], http=FakeHttp(), send=send)
        return {m["to"]: m for m in sent if " wrapped · " in m["subject"]}
    wr = wrap_run("2027-01-08T13:30:00Z")
    check("wrapped: on the 8th of January every account with a portfolio gets its 2026 (the owner's from the main portfolio), to its own address",
          sorted(wr) == ["friend2@example.com", "friend@example.com", "hist2@example.com", "owner@example.com"] and all(": your 2026 wrapped · " in m["subject"] for m in wr.values())
          and "2026 IN NUMBERS" in wr["owner@example.com"]["text"], json.dumps({k: v["subject"] for k, v in wr.items()}))
    check("wrapped: ranked among friends (percentages only), the account without friends without a ranking",
          "YOU AND YOUR FRIENDS" in wr["friend@example.com"]["text"] and "Friend Portfolio" in wr["friend@example.com"]["text"] and "YOU AND YOUR FRIENDS" not in wr["hist2@example.com"]["text"])
    check("wrapped: once a year", wrap_run("2027-01-08T17:30:00Z") == {} and wrap_run("2027-02-02T13:30:00Z") == {})

    # ---- the morning brief (run_morning.py): the owner's main portfolio, and the accounts that switched it on ----
    import run_morning, jobs_common as jc
    opkg = DB[f"mail/{OUID}"]["pkg"]
    json.dump({"portfolioId": "khaled", "name": "Main", "siteRepo": "x/y", "siteFolder": "p/khaled"}, open(os.path.join(eng, "config.json"), "w"))
    real_record = jc.record_job
    jc.record_job = lambda ctx, section, updates, message: (lambda st: (st.setdefault(section, {}).update(updates), jc.save_jobs_state(ctx, st)))(jc.jobs_state(ctx))
    owner_sent = []
    def morning(now):
        sent.clear(); owner_sent.clear()
        run_morning.main(["--engine", eng, "--code", code, "--now", now], http=FakeHttp(), send=send, send_owner=lambda ctx, s, t, h: owner_sent.append((s, t)) or "sent")
        return [m for m in sent if "morning brief" in m["subject"]]
    try:
        DB[f"mail/{UID}"]["pkg"] = seal_mail({"uid": UID, "email": "friend@example.com", "refresh": "RT1", "pk8": pk8, "prefs": {"alerts": True, "morning": True}})
        DB[f"mail/{OUID}"]["pkg"] = seal_mail({"uid": OUID, "email": "owner@example.com", "refresh": "RT3", "pk8": opk8, "prefs": {"shareMain": True, "morning": True}})
        ms = morning("2026-09-24T06:00:00Z")      # Thursday 9:00 Cairo
        check("morning brief: the owner's main portfolio, to the owner's address (through the owner's mail lock)",
              len(owner_sent) == 1 and owner_sent[0][0].startswith("Demo Portfolio: morning brief") and "YOUR HOLDINGS" in owner_sent[0][1], json.dumps([x[0] for x in owner_sent]))
        check("morning brief: only the account that switched it on gets one, to its own address (not the owner's sign-in account, not the others)",
              [m["to"] for m in ms] == ["friend@example.com"] and "morning brief" in ms[0]["subject"] and "YOUR HOLDINGS" in ms[0]["text"], json.dumps([[m["to"], m["subject"]] for m in ms]))
        check("morning brief: once a day", morning("2026-09-24T06:20:00Z") == [] and owner_sent == [])
        check("morning brief: nothing on a Friday", morning("2026-09-25T06:00:00Z") == [] and owner_sent == [])
        check("morning brief: the next session day, again", len(morning("2026-09-27T06:00:00Z")) == 1 and len(owner_sent) == 1)
    finally:
        jc.record_job = real_record
        DB[f"mail/{UID}"]["pkg"], DB[f"mail/{OUID}"]["pkg"] = upkg, opkg

    # ---- moving the main portfolio into the owner's account: copy, shadow (no emails), then live ----
    import migrate_main
    os.environ["GMAIL_ADDRESS"], os.environ["GMAIL_APP_PASSWORD"] = "owner.gmail@example.com", "abcdefghijklmnop"
    cfg_path = os.path.join(eng, "config.json")
    cfg0 = open(cfg_path).read() if os.path.exists(cfg_path) else None
    try:
        before = set(k for k in DB if k.startswith(f"users/{OUID}/docs/"))
        rc_m = migrate_main.main(["--engine", eng, "--code", code, "--mode", "copy"], http=FakeHttp())
        oset = adoc(OUID, oacct, "portfolio/settings") or {}
        check("migrate: the main portfolio is copied into the owner's account in one go (market data stays shared)",
              rc_m == 0 and oset.get("migratedFrom") and oset.get("inception") and adoc(OUID, oacct, "ledger/y2026") is not None
              and not any("/market__" in k or "/history__" in k for k in DB if k.startswith(f"users/{OUID}/docs/"))
              and (adoc(OUID, oacct, "sync/gmail") or {}).get("address") == "owner.gmail@example.com", json.dumps(sorted(k.rsplit('/', 1)[1] for k in DB if k.startswith(f'users/{OUID}/docs/'))))
        check("migrate: a second copy is refused", migrate_main.main(["--engine", eng, "--code", code, "--mode", "copy"], http=FakeHttp()) == 1)
        # the main portfolio's month-end files move too: sealed again to the account's key, in its own folder on the site
        mexp = os.path.join(code, "p", "khaled", "exports")
        os.makedirs(mexp, exist_ok=True)
        mkeys = os.path.join(code, "p", "khaled", "keys.json")
        open(os.path.join(mexp, "Main-Aug-26.xlsx.enc.json"), "wb").write(store.seal(mkeys, b"PK-main-workbook", "Main-Aug-26.xlsx"))
        json.dump([{"month": "2026-08", "name": "Main-Aug-26.xlsx", "file": "exports/Main-Aug-26.xlsx.enc.json"}], open(os.path.join(mexp, "index.json"), "w"))
        rc_f = migrate_main.main(["--engine", eng, "--code", code, "--mode", "files"], http=FakeHttp())
        chk2 = os.path.join(tmp, "site-check2")
        subprocess.run(["git", "clone", "-q", os.environ["SITE_REMOTE"], chk2], check=True, capture_output=True)
        ofd = os.path.join(chk2, ram.files_dir(OUID))
        oidx = json.loads(store.unseal(oacct, open(os.path.join(ofd, "index.enc.json"), "rb").read())) if os.path.exists(os.path.join(ofd, "index.enc.json")) else []
        check("migrate: the main portfolio's month-end files are in the owner account's folder, opened with the account's key",
              rc_f == 0 and oidx and oidx[0]["month"] == "2026-08" and store.unseal(oacct, open(os.path.join(chk2, oidx[0]["file"]), "rb").read()) == b"PK-main-workbook", json.dumps(oidx))
        sent.clear(); fetches.clear()
        ram.main(["--engine", eng, "--code", code, "--now", "2026-09-29T19:30:00Z"], http=FakeHttp(), send=send)
        check("shadow: the owner's account reads the owner's Gmail and saves the import, but nothing is emailed to the owner",
              any(f["addr"] == "owner.gmail@example.com" for f in fetches) and not [m for m in sent if m["to"] == "owner@example.com"]
              and ((adoc(OUID, oacct, "sync/state") or {}).get("seen") or {}).get("18a0b0c0d0e0f001"), json.dumps([[m["to"], m["subject"]] for m in sent]))
        check("shadow: the one login stays (signing in still opens the main portfolio)", "mainKey" in PROFILES.get(OUID, {}))
        cfg = json.load(open(cfg_path)) if cfg0 else {"portfolioId": "khaled", "name": "Main", "siteRepo": "x/y", "siteFolder": "p/khaled"}
        json.dump({**cfg, "movedToAccount": True}, open(cfg_path, "w"))
        sent.clear()
        ram.main(["--engine", eng, "--code", code, "--now", "2026-09-24T19:30:00Z", "--weekly"], http=FakeHttp(), send=send)
        mine = [m["subject"] for m in sent if m["to"] == "owner@example.com"]
        check("live: the owner's account is an ordinary account now: its own weekly summary, and the one login is removed",
              any("week" in x.lower() for x in mine) and "mainKey" not in PROFILES.get(OUID, {}), json.dumps(mine))
        owner_sent.clear(); sent.clear()
        run_morning.main(["--engine", eng, "--code", code, "--now", "2026-10-04T06:00:00Z"], http=FakeHttp(), send=send, send_owner=lambda ctx, s, t, h: owner_sent.append((s, t)) or "sent")
        check("live: the morning brief comes from the account, not from the engine",
              owner_sent == [] and [m["to"] for m in sent if "morning brief" in m["subject"]].count("owner@example.com") == 1, json.dumps([[m["to"], m["subject"]] for m in sent]))
    finally:
        if cfg0 is None:
            os.path.exists(cfg_path) and os.remove(cfg_path)
        else:
            open(cfg_path, "w").write(cfg0)

    # ---- a friendship made by @username: the asker's side has no email until it is accepted ----
    class HandleHttp:
        DOCS = {"handles/omar": {"uid": "Uomar", "pub": "PUB-omar"}, "directory/omar@example.com": {"uid": "Uomar", "pub": "PUB-omar"}}
        def json(self, method, url, headers=None, body=None):
            d = self.DOCS.get(urllib.parse.unquote(url.split("/documents/", 1)[1]))
            return (200, {"fields": {k: {"stringValue": v} for k, v in d.items()}}) if d else (404, {})
    hh = HandleHttp()
    check("@username link: a request sent to @omar is confirmed by handles/omar (same account, same key)",
          ram.confirmed_friend(hh, "T", {"uid": "Uomar", "pub": "PUB-omar", "handle": "omar"}) is True)
    check("@username link: someone else's key under @omar is not confirmed, nor an unknown username",
          ram.confirmed_friend(hh, "T", {"uid": "Uomar", "pub": "PUB-evil", "handle": "@Omar"}) is False
          and ram.confirmed_friend(hh, "T", {"uid": "Uomar", "pub": "PUB-omar", "handle": "nobody"}) is False)
    check("@username link: once accepted it carries the friend's email, and the email decides",
          ram.confirmed_friend(hh, "T", {"uid": "Uomar", "pub": "PUB-omar", "handle": "renamed", "email": "omar@example.com"}) is True)
finally:
    shutil.rmtree(tmp, ignore_errors=True)
print("ALL PASS" if not fails else f"{fails} FAILED")
sys.exit(1 if fails else 0)
