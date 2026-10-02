#!/usr/bin/env python3
"""Month-end Excel workbook from the JSON that excel.js writes.
   python3 excel.py <data.json> <out.xlsx>      (also writes <out.xlsx>.b64 for emailing)
   Sheets: Summary, Monthly, Holdings, Ledger, Closed trades, Income, Attribution, Marks & inputs, Monthly (formulas).
   Values are the page's figures everywhere (openpyxl cannot store a cached value with a formula, so a formula-only sheet shows
   blanks in phone previews); the last sheet repeats the Monthly table with live formulas so the return chain stays auditable.
   Text from the data that starts with = + - @ is always stored as text, never as a formula."""
import sys, json, base64, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

d = json.load(open(sys.argv[1]))
wb = Workbook()
INK = "1D1D1F"; ACC = "0071E3"; RULE = "D2D2D7"; INPUT = "2256C7"   # the site's palette; typed-input cells keep the usual finance blue
H = Font(bold=True, color="FFFFFF", size=10); HF = PatternFill("solid", fgColor=ACC)
T = Font(bold=True, size=14, color=INK); S = Font(size=9, color="6E6E73"); B = Font(bold=True)
thin = Side(style="thin", color=RULE); bottom = Border(bottom=thin)
PCT = "+0.0%;-0.0%;0.0%"; PCT2 = "+0.00%;-0.00%;0.00%"; PCTU = "0.0%"; EGP = "#,##0;(#,##0);-"; EGP2 = "#,##0.00;(#,##0.00);-"; NUM = "0.00"; DT = "yyyy-mm-dd"
MON = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
lbl = lambda m: f"{MON[int(m[5:7]) - 1]}-{m[2:4]}"
FORMULA_START = ("=", "+", "-", "@")

def put(ws, r, c, v=None, formula=False):
    """ws.cell() that stores data text starting with = + - @ as text (no formula injection); formula=True keeps real formulas."""
    cell = ws.cell(r, c, v)
    if not formula and isinstance(v, str) and v.startswith(FORMULA_START):
        cell.value = v; cell.data_type = "s"
    return cell

def sheet(title, header, rows, fmts=None, widths=None, freeze="A2", note=None, start=1, formulas=False):
    ws = wb.create_sheet(title) if wb.active.title != "Sheet" else wb.active
    ws.title = title
    r0 = start
    if note:
        put(ws, r0, 1, note).font = S; r0 += 2
    for j, h in enumerate(header, 1):
        c = put(ws, r0, j, h); c.font = H; c.fill = HF; c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[r0].height = 30
    for i, row in enumerate(rows, r0 + 1):
        for j, v in enumerate(row, 1):
            c = put(ws, i, j, v, formulas)
            if fmts and fmts[j - 1]: c.number_format = fmts[j - 1]
            if isinstance(v, (int, float)) and not isinstance(v, bool): c.alignment = Alignment(horizontal="right")
    for j, w in enumerate(widths or [14] * len(header), 1): ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(r0 + 1, 1)
    ws.sheet_view.showGridLines = False
    return ws, r0

# ---------- Summary ----------
ws = wb.active; ws.title = "Summary"; ws.sheet_view.showGridLines = False
put(ws, 1, 1, d["name"]).font = T
ws["A2"] = f"Month-end report · {lbl(d['month'])} · benchmark {d['benchmark']} · EGP · inception {lbl(d['inception'])} · generated {d['generated']}"; ws["A2"].font = S
r = 4
groups = [("Value", ["Portfolio value (EGP)", "Cash (EGP)", "Securities (EGP)", "Difference: statement value vs closing prices (EGP)"]),
          ("Returns", ["Month return", "Benchmark month return", "Month alpha", "Since inception TWR", "Since inception benchmark", "Alpha since inception", "Annualized TWR", "Money-weighted return (XIRR, annual)", "Return in USD", "Real return (after CPI)",
                       "Cash benchmark (CBE policy rate) over the period", "Ahead of cash", "Index dividend yield (estimate)", "Index return with dividends (estimate)", "Return before trading costs"]),
          ("Risk", ["Monthly volatility", "Annualized volatility", "Sharpe ratio", "Sortino ratio", "Calmar ratio", "Beta vs benchmark", "Jensen's alpha (annual)", "Correlation", "Tracking error (annual)", "Upside capture", "Downside capture", "Max drawdown (month-end)", "Positive months", "Months in period", "Months beating benchmark"]),
          ("Capital", ["Opening value before inception (EGP)", "Deposits since inception (EGP)", "Withdrawals since inception (EGP)", "Investment gain since inception (EGP)", "Dividends received since inception (EGP)", "Realized trading P/L (EGP)", "Trading costs (EGP)", "Trading costs as % of value traded", "Risk-free rate used"])]
# returns are signed with 2 decimals (as on the Monthly sheet); unsigned quantities (rates, volatility, capture ratios) plain 0.0%
ret_keys = {"Month return", "Benchmark month return", "Month alpha", "Since inception TWR", "Since inception benchmark", "Alpha since inception", "Annualized TWR", "Money-weighted return (XIRR, annual)", "Return in USD", "Real return (after CPI)", "Max drawdown (month-end)",
            "Cash benchmark (CBE policy rate) over the period", "Ahead of cash", "Index return with dividends (estimate)", "Return before trading costs", "Jensen's alpha (annual)"}
unsigned_keys = {"Monthly volatility", "Annualized volatility", "Tracking error (annual)", "Upside capture", "Downside capture", "Risk-free rate used", "Index dividend yield (estimate)"}
small_pct_keys = {"Trading costs as % of value traded"}  # ~0.1-0.5%: two decimals
labels = d.get("labels") or {}
for g, keys in groups:
    put(ws, r, 1, g.upper()).font = Font(bold=True, size=9, color="6E6E73"); r += 1
    for k in keys:
        v = d["summary"].get(k)
        put(ws, r, 1, labels.get(k, k)).border = bottom
        c = put(ws, r, 2, v if v is not None else "—"); c.border = bottom; c.alignment = Alignment(horizontal="right")
        c.number_format = PCT2 if k in ret_keys else PCTU if k in unsigned_keys else "0.00%" if k in small_pct_keys else EGP if "(EGP)" in k else NUM if isinstance(v, float) else "0"
        r += 1
    r += 1
put(ws, r, 1, "TRAILING RETURNS").font = Font(bold=True, size=9, color="6E6E73"); r += 1
for j, h in enumerate(["Period", "Portfolio", "Benchmark", "Difference"], 1):
    c = put(ws, r, j, h); c.font = H; c.fill = HF
r += 1
for t in d["trailing"]:
    put(ws, r, 1, t["period"])
    for j, k in enumerate(["portfolio", "benchmark", "difference"], 2):
        c = put(ws, r, j, t[k] if t.get(k) is not None else "—"); c.number_format = PCT2; c.alignment = Alignment(horizontal="right")
    r += 1
ws.column_dimensions["A"].width = 42; ws.column_dimensions["B"].width = 18; ws.column_dimensions["C"].width = 14; ws.column_dimensions["D"].width = 14

# ---------- Monthly (values) ----------
# Columns: A Month, B Opening, C Deposits, D Withdrawals, E Net flow, F Day-weighted flow, G Dividends, H Cash, I Securities,
#          J Month-end value, K Return, L Benchmark, M Alpha, N Cumulative, O Cumulative benchmark, P Drawdown, Q close, R USD/EGP, S CPI,
#          T USD return, U Real return, V Cash return (CBE policy rate), W Trades, X Trading costs, Y source
hdr = ["Month", "Opening value", "Deposits", "Withdrawals", "Net flow", "Day-weighted flow", "Dividends", "Cash", "Securities", "Month-end value", "Return", "Benchmark", "Alpha", "Cumulative", "Cumulative benchmark", "Drawdown", "EGX30 Capped close", "USD/EGP", "CPI MoM", "USD return", "Real return", "Cash return", "Trades", "Trading costs", "Month-end source"]
fm = [None, EGP, EGP, EGP, EGP, EGP2, EGP, EGP2, EGP2, EGP2, PCT2, PCT2, PCT2, PCT, PCT, PCT, "#,##0.0", "0.00", PCT2, PCT, PCT, PCT2, "0", EGP2, None]
mwidths = [9, 14, 12, 12, 12, 13, 11, 13, 14, 15, 10, 11, 10, 11, 12, 10, 13, 9, 9, 10, 10, 10, 7, 12, 22]
rows = [[lbl(m["month"]), m["opening"], m["deposits"], m["withdrawals"], m["netFlow"], m.get("weightedFlow", m["deposits"] - m["withdrawals"]), m["dividends"], m["cash"], m["securities"], m["value"],
         m["ret"], m["bench"], m["alpha"], m["cum"], m["cumBench"], m["dd"], m["benchClose"], m["usdegp"], m["cpi"], m["usdRet"], m["realRet"], m.get("cashRet"), m["trades"], m.get("tradingCost"), m["source"]] for m in d["monthly"]]
sheet("Monthly", hdr, rows, fm, mwidths,
      note="Return = Modified Dietz: (month-end value − opening − net flow) ÷ (opening + day-weighted net flow), where each deposit or withdrawal is weighted by the share of the month it was invested (a flow on the 1st counts fully); months are chain-linked. Cash return = the CBE policy rate for the month ((1 + annual rate)^(1/12) − 1; the settings' risk-free rate where no CBE rate is recorded). Trading costs = net amount vs price × shares on stock trades. All figures are the page's values; the last sheet, 'Monthly (formulas)', repeats this table with live formulas.")

# ---------- Holdings ----------
diff = d["summary"].get("Difference: statement value vs closing prices (EGP)") or 0
rows = [[h["symbol"], h["name"], h["sector"], h["shares"], h["avgCost"], h["price"], h["priceSource"], h["cost"], h["mv"], h["unreal"], h["ret"], h["weight"], h["dividends"]] for h in d["holdings"]]
rows.append(["", "Cash", "Cash", None, None, None, "", None, d["holdingsCash"], None, None, (d["holdingsCash"] / d["holdingsTotal"]) if d["holdingsTotal"] else 0, None])
rows.append(["", "TOTAL", "", None, None, None, "", sum(h["cost"] for h in d["holdings"]), d["holdingsTotal"], sum((h["unreal"] or 0) for h in d["holdings"]), None, 1, sum(h["dividends"] for h in d["holdings"])])
ws, r0 = sheet("Holdings", ["Symbol", "Company", "Sector", "Shares", "Avg cost", "Price", "Price source", "Cost basis", "Market value", "Unrealized P/L", "Return", "Weight", "Dividends received"], rows,
      [None, None, None, "#,##0.####", "0.000", "0.00", None, EGP, EGP, EGP, PCT, "0.0%", EGP], [10, 38, 22, 12, 10, 9, 11, 14, 14, 14, 9, 8, 13],
      note=f"Positions at {d['month']} month-end: stocks at closing prices, funds at their last traded NAV (gold funds moved with the 24K gold price). Average-cost basis. "
           + (f"The Summary's portfolio value is the month-end mark (Thndr's statement figure); valuing the same shares at closing prices gives this total, {abs(diff):,.0f} EGP {'lower' if diff > 0 else 'higher'} (Summary: 'Difference: statement value vs closing prices')." if abs(diff) >= 0.5 else "This total equals the Summary's portfolio value."))
for c in ws[r0 + len(rows)]: c.font = B  # the TOTAL row

# ---------- Ledger ----------
rows = [[datetime.date.fromisoformat(t["date"]), t["type"], t["asset"], t["symbol"], t["shares"], t["price"], t["amount"], t["basisSold"], t["realized"], t["account"], t["source"], t["note"]] for t in d["ledger"]]
sheet("Ledger", ["Date", "Type", "Asset", "Symbol", "Shares", "Price", "Net amount", "Cost basis sold", "Realized P/L", "Account", "Source", "Note"], rows,
      [DT, None, None, None, "#,##0.####", "0.####", EGP2, EGP2, EGP2, None, None, None], [11, 10, 38, 9, 12, 10, 14, 14, 13, 8, 16, 30],
      note="Every transaction as reconciled with Thndr statements. Buys, withdrawals and fees are negative; sells, deposits, dividends and rebates are positive.")

# ---------- Closed trades ----------
trip_n = {}
for c in d["closed"]: trip_n[c["name"]] = trip_n.get(c["name"], 0) + 1
rows = [[c["name"] + (f" ({c['trip']})" if trip_n[c["name"]] > 1 and c.get("trip") else ""), c["symbol"], datetime.date.fromisoformat(c["firstBuy"]), datetime.date.fromisoformat(c["lastSell"]), c["holdDays"], c["buyCost"], c["proceeds"], c["dividends"], c["total"], c["roi"], c["outcome"]] for c in d["closed"]]
sheet("Closed trades", ["Asset", "Symbol", "First buy", "Last sell", "Days held", "Buy cost", "Sell proceeds", "Dividends", "Total P/L", "ROI", "Outcome"], rows,
      [None, None, DT, DT, "0", EGP, EGP, EGP, EGP, PCT, None], [38, 10, 11, 11, 9, 13, 13, 11, 13, 9, 9],
      note="Round trips fully exited by month-end (a stock sold out and bought again later is a new trip, numbered 'Name (2)'). P/L covers every lot of the trip, including its dividends.")

# ---------- Income ----------
rows = [[Y["year"]] + Y["div"] + [Y["divT"], Y["rebT"], Y["feeT"], Y["net"]] for Y in d["incomeYears"]]
ws, r0 = sheet("Income", ["Year"] + MON + ["Dividends", "Rebates", "Fees", "Net income"], rows, [None] + [EGP] * 16, [8] + [9] * 12 + [12, 11, 10, 12],
      note="Dividends by month (EGP). Rebates are Thndr commission kickbacks; fees are subscriptions, custody and transfers.")
r = r0 + len(rows) + 3
put(ws, r, 1, "DIVIDENDS BY STOCK").font = Font(bold=True, size=9, color="6E6E73"); r += 1
for j, h in enumerate(["Stock", "Symbol", "Payments", "Total received", "Last paid"], 1):
    c = put(ws, r, j, h); c.font = H; c.fill = HF
r += 1
for x in d["incomeByStock"]:
    put(ws, r, 1, x["name"]); put(ws, r, 2, x["symbol"]); put(ws, r, 3, x["payments"]); c = put(ws, r, 4, x["total"]); c.number_format = EGP
    put(ws, r, 5, datetime.date.fromisoformat(x["last"]) if x["last"] else "").number_format = DT; r += 1
ws.column_dimensions["A"].width = 30

# ---------- Attribution ----------
A = d.get("attribution")
if A:
    rows = [[s["sector"], s["wp"], s["wb"], s["wp"] - s["wb"], s["rp"], s["rb"], s["alloc"], s["sel"], s["total"]] for s in A["sectors"]]
    rows.append(["TOTAL", 1, 1, 0, None, None, A["allocation"], A["selection"], A["allocation"] + A["selection"]])
    ws, r0 = sheet("Attribution", ["Sector", "Portfolio weight", "Index weight", "Active weight", "Portfolio return", "Index return", "Allocation", "Selection", "Total"], rows,
          [None, "0.0%", "0.0%", "+0.0%;-0.0%;0.0%", PCT, PCT, PCT2, PCT2, PCT2], [26, 12, 12, 12, 12, 12, 11, 11, 11],
          note=f"Brinson-Fachler vs a free-float model of EGX30 Capped, since inception to {lbl(d['month'])}, Carino-linked. Active return {A['active']:+.2%} = allocation {A['allocation']:+.2%} + selection {A['selection']:+.2%} + trading and other {A['trading']:+.2%} + index model gap {A['replication']:+.2%}. " + (A.get("tradingNote") or "") + ".")
    for c in ws[r0 + len(rows)]: c.font = B

# ---------- Marks & inputs ----------
# the CPI / USD-EGP source columns only appear when at least one month names a source
mcols = [("Month", lambda m: lbl(m["month"]), None, 9), ("Cash", lambda m: m.get("cash"), EGP2, 14), ("Securities", lambda m: m.get("securities"), EGP2, 15),
         ("EGX30 Capped close", lambda m: m.get("benchClose"), "#,##0.0", 14), ("EGX30 Capped % (typed)", lambda m: m.get("benchReturn"), PCT2, 12),
         ("USD/EGP", lambda m: m.get("usdegp"), "0.00", 9), ("CPI MoM", lambda m: m.get("cpi"), PCT2, 9), ("Month-end source", lambda m: m.get("source", "typed"), None, 15)]
if any(isinstance(m.get("cashRate"), (int, float)) for m in d["marks"]): mcols.append(("CBE policy rate (annual)", lambda m: m.get("cashRate"), "0.00%", 12))
for key, title in (("cpiSource", "CPI source"), ("usdegpSource", "USD/EGP source")):
    if any(m.get(key) for m in d["marks"]): mcols.append((title, (lambda k: lambda m: m.get(k, ""))(key), None, 30))
mcols.append(("Note", lambda m: m.get("note", ""), None, 30))
rows = [[f(m) for _, f, _, _ in mcols] for m in d["marks"]]
ws, r0 = sheet("Marks & inputs", [t for t, _, _, _ in mcols], rows, [x for _, _, x, _ in mcols], [w for _, _, _, w in mcols],
      note="Month-end inputs behind every figure. Blue = values entered or imported (spreadsheet convention for hard-coded inputs).")
for row in ws.iter_rows(min_row=r0 + 1, max_row=r0 + len(rows), min_col=2, max_col=7):
    for c in row: c.font = Font(color=INPUT)
r = r0 + len(rows) + 3
put(ws, r, 1, "SETTINGS").font = Font(bold=True, size=9, color="6E6E73"); r += 1
for k in ["name", "inception", "openingValue", "benchCloseStart", "fxStart", "riskFree", "cash", "cashDate", "cashSource", "openThreshold", "staleDays"]:
    put(ws, r, 1, k); c = put(ws, r, 2, d["settings"].get(k)); c.font = Font(color=INPUT); r += 1
r += 2
put(ws, r, 1, "ASSETS").font = Font(bold=True, size=9, color="6E6E73"); r += 1
for j, h in enumerate(["Asset", "Symbol", "Sector", "Target", "Stop loss", "Fund"], 1):
    c = put(ws, r, j, h); c.font = H; c.fill = HF
r += 1
for a in d["assets"]:
    for j, v in enumerate([a["name"], a["symbol"], a["sector"], a["target"], a["stop"], "yes" if a["fund"] else ""], 1): put(ws, r, j, v)
    r += 1

# ---------- Monthly (formulas): the same table with the return chain as live formulas ----------
rows = []
for i, m in enumerate(d["monthly"]):
    R = 4 + i  # note row, blank, header, then data
    rows.append([lbl(m["month"]), m["opening"], m["deposits"], m["withdrawals"], f"=C{R}-D{R}", m.get("weightedFlow", m["deposits"] - m["withdrawals"]), m["dividends"], m["cash"], m["securities"], f"=H{R}+I{R}",
                 f"=(J{R}-B{R}-E{R})/(B{R}+F{R})", m["bench"], f"=K{R}-L{R}", f"=(1+K{R})*{'(1+N' + str(R - 1) + ')' if i else '1'}-1", f"=(1+L{R})*{'(1+O' + str(R - 1) + ')' if i else '1'}-1",
                 f"=(1+N{R})/MAX(1,1+MAX($N$4:N{R}))-1", m["benchClose"], m["usdegp"], m["cpi"], m["usdRet"], m["realRet"], m.get("cashRet"), m["trades"], m.get("tradingCost"), m["source"]])
sheet("Monthly (formulas)", hdr, rows, fm, mwidths, formulas=True,
      note="Recalculates in Excel/Numbers/Sheets; phone previews show the Monthly sheet. Columns E, J, K, M, N, O, P are live formulas (Modified Dietz, chain-linked); the rest are the page's figures.")

wb.save(sys.argv[2])
b = open(sys.argv[2], "rb").read()
open(sys.argv[2] + ".b64", "w").write(base64.b64encode(b).decode())
print(json.dumps({"ok": True, "file": sys.argv[2], "bytes": len(b), "sheets": wb.sheetnames}))
